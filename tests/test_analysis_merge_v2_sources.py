# -*- coding: utf-8 -*-
"""분석본 결합 v2 — 교과서/EBS 소스 확장(W28-확장) 오프라인 테스트.

스캔(합성 트리) + 범위 유도(실물 전례 정합) + 파일명 축약 + verify doc 게이트 +
JSX/폼 계약. InDesign 비의존.
"""
import json
import os
import sys

import pytest

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(TESTS_DIR)
sys.path.insert(0, REPO)

from engines.analysis_merge_v2 import sources as S            # noqa: E402
from engines.analysis_merge_v2.engine import AnalysisMergeV2Engine  # noqa: E402


@pytest.fixture
def engine():
    return AnalysisMergeV2Engine()


# ── 합성 소스 트리 ───────────────────────────────────────────────────────────

@pytest.fixture
def tb_tree(tmp_path):
    """실물 구조 재현: 과목/출판사/과 + 학교 변형본 + 잡파일/dot-dir."""
    base = tmp_path / "교과서 작업"
    lesson = base / "공통영어1" / "공통영어1 비상(홍)" / "공통영어1 비상(홍) 3과"
    lesson.mkdir(parents=True)
    (lesson / "공통영어1 비상(홍) 3과.indd").write_bytes(b"x")
    (lesson / "공통영어1 비상(홍) 3과 세종1.indd").write_bytes(b"x")
    (lesson / "공통영어1비상(홍)Lesson3_어법수정_문제지.docx").write_bytes(b"x")
    (lesson / "~공통영어1 비상(홍) 3과.idlk").write_bytes(b"x")
    sub = lesson / "공통영어1 비상(홍) 3과 워크북"
    sub.mkdir()
    (sub / "속지.indd").write_bytes(b"x")          # 하위 폴더 — 내리지 않음
    # indd 없는 과 (미제작) — 목록 제외
    (base / "공통영어1" / "공통영어1 비상(홍)" / "공통영어1 비상(홍) 1과").mkdir()
    # 제외 대상
    (base / "이전 파일(15개정)").mkdir()
    (base / ".rsync-partial").mkdir()
    return {"textbook_base": str(base), "ebs_base": str(tmp_path / "없음")}


@pytest.fixture
def ebs_tree(tmp_path):
    base = tmp_path / "EBS 작업"
    # ① N강 폴더 + 파일명 표기 변형 (실측: "올림포스 영어독해 기본2 15강.indd")
    u15 = base / "올림포스 독해기본2" / "올림포스 독해기본2 15강"
    u15.mkdir(parents=True)
    (u15 / "올림포스 영어독해 기본2 15강.indd").write_bytes(b"x")
    u1 = base / "올림포스 독해기본2" / "올림포스 독해기본2 1강"
    u1.mkdir()
    (u1 / "올림포스 독해기본2 1강.indd").write_bytes(b"x")
    (base / "올림포스 독해기본2" / "올독기2 꼼꼼변형 한글").mkdir()   # 강 아님
    # 범위 묶음 폴더
    u12 = base / "올림포스 독해기본1" / "올림포스 독해기본1 1~2강"
    u12.mkdir(parents=True)
    (u12 / "올림포스 1~2강.indd").write_bytes(b"x")
    # ② Unit 폴더 (Reading Power)
    rp = base / "Reading Power 유형편(기본)" / "Unit 13 Example, 1, 2"
    rp.mkdir(parents=True)
    (rp / "Reading Power 유형편(기본) Unit 13 Ex, 1, 2.indd").write_bytes(b"x")
    (rp / "Reading Power 유형편(기본) Unit 13 Ex, 1, 2 세종1.indd").write_bytes(b"x")
    # ③ 교재 직속 파일 (수능특강)
    st = base / "2026 수능특강 영어"
    st.mkdir()
    (st / "2026 수능특강 영어(12-15강) 명여버전(완료).indd").write_bytes(b"x")
    (st / "2026 수능특강 영어(Test 2) 함원장버전.indd").write_bytes(b"x")
    return {"textbook_base": str(tmp_path / "없음"), "ebs_base": str(base)}


# ── 교과서 스캔 ──────────────────────────────────────────────────────────────

def test_tb_subjects_exclude_and_prune(tb_tree):
    assert S.list_textbook_subjects(tb_tree) == ["공통영어1"]


def test_tb_lessons_school_variants(tb_tree):
    lessons = S.list_textbook_lessons(tb_tree, "공통영어1", "공통영어1 비상(홍)")
    assert len(lessons) == 1                      # indd 없는 1과는 제외
    e = lessons[0]
    assert e["folder"] == "공통영어1 비상(홍) 3과"
    assert e["label"] == "3과"
    # 기본본 먼저, 학교 변형본 구분 / 잡파일·락파일·하위폴더 미포함
    assert e["files"] == [
        {"name": "공통영어1 비상(홍) 3과.indd", "school": None},
        {"name": "공통영어1 비상(홍) 3과 세종1.indd", "school": "세종1"},
    ]


def test_tb_resolve_roundtrip(tb_tree):
    path, label, school = S.resolve_textbook_file(
        tb_tree, "공통영어1", "공통영어1 비상(홍)",
        "공통영어1 비상(홍) 3과", "공통영어1 비상(홍) 3과 세종1.indd")
    assert path.endswith("공통영어1 비상(홍) 3과 세종1.indd")
    assert label == "3과" and school == "세종1"
    with pytest.raises(ValueError):
        S.resolve_textbook_file(tb_tree, "공통영어1", "공통영어1 비상(홍)",
                                "공통영어1 비상(홍) 3과", "없는파일.indd")


# ── EBS 스캔 ─────────────────────────────────────────────────────────────────

def test_ebs_unit_variant_filename_and_boundary(ebs_tree):
    units = S.list_ebs_units(ebs_tree, "올림포스 독해기본2")
    by_label = {u["label"]: u for u in units}
    assert set(by_label) == {"1강", "15강"}
    # 파일명 표기 변형 매칭 + "1강"↛"15강" 오매칭 없음
    assert by_label["15강"]["files"] == [
        {"name": "올림포스 영어독해 기본2 15강.indd", "school": None}]
    assert by_label["1강"]["files"] == [
        {"name": "올림포스 독해기본2 1강.indd", "school": None}]


def test_ebs_range_folder_label(ebs_tree):
    units = S.list_ebs_units(ebs_tree, "올림포스 독해기본1")
    assert units[0]["label"] == "1~2강"
    assert units[0]["files"][0]["name"] == "올림포스 1~2강.indd"


def test_ebs_unit_style_folders(ebs_tree):
    units = S.list_ebs_units(ebs_tree, "Reading Power 유형편(기본)")
    assert len(units) == 1 and units[0]["label"] == "Unit 13"
    assert units[0]["files"][0]["school"] is None
    assert units[0]["files"][1]["school"] == "세종1"


def test_ebs_loose_files_suneungteukgang(ebs_tree):
    units = S.list_ebs_units(ebs_tree, "2026 수능특강 영어")
    by_label = {u["label"]: u for u in units}
    assert set(by_label) == {"12-15강", "Test2"}
    assert by_label["12-15강"]["folder"] == ""     # 교재 직속 표시
    assert by_label["12-15강"]["files"][0]["school"] == "명여버전 완료"


def test_ebs_resolve_loose_file(ebs_tree):
    path, label, school = S.resolve_ebs_file(
        ebs_tree, "2026 수능특강 영어", "",
        "2026 수능특강 영어(Test 2) 함원장버전.indd")
    assert path.endswith("함원장버전.indd") and label == "Test2"
    assert school == "함원장버전"


def test_ebs_nested_and_book_level_analysis_folders(tmp_path):
    """실측 2026-09-14 — 전부 '강 없음 (미지원 구조)'로 막혀 있던 배치."""
    base = tmp_path / "EBS 작업"
    lt = base / "수능특강라이트 영어(개정판)" / "13강"
    (lt / "분석본").mkdir(parents=True)
    (lt / "분석본" / "13강_분석본.indd").write_bytes(b"x")
    (lt / "분석본" / "13강 분석본_이의고1.indd").write_bytes(b"x")
    (lt / "리테").mkdir()
    (lt / "리테" / "수능특강 Light 영어 13강 리테.indd").write_bytes(b"x")
    rp = base / "Reading Power 유형편(완성)" / "분석본"
    rp.mkdir(parents=True)
    (rp / "Reading Power 유형편(완성) Unit 18.indd").write_bytes(b"x")
    st = base / "2027 수능특강 영어" / "분석본"
    for sub in ("유형편", "내지", "TEST 1"):
        (st / sub).mkdir(parents=True)
    (st / "유형편" / "고3_수능특강_분석본_03강.indd").write_bytes(b"x")
    (st / "내지" / "고3_분석본_내지 복사본.indd").write_bytes(b"x")
    (st / "TEST 1" / "고3_분석본_TEST 1_17~20번.indd").write_bytes(b"x")
    cfg = {"textbook_base": "", "ebs_base": str(base)}

    u = S.list_ebs_units(cfg, "수능특강라이트 영어(개정판)")
    assert [(x["label"], x["folder"]) for x in u] == [("13강", "13강/분석본")]
    assert u[0]["files"] == [
        {"name": "13강_분석본.indd", "school": None},
        {"name": "13강 분석본_이의고1.indd", "school": "이의고1"}]
    path, label, school = S.resolve_ebs_file(
        cfg, "수능특강라이트 영어(개정판)", "13강/분석본", "13강 분석본_이의고1.indd")
    assert path.endswith("13강 분석본_이의고1.indd")
    assert label == "13강" and school == "이의고1"
    r = S.derive_range_ebs("수능특강라이트 영어(개정판)", "13강", "13강_분석본.indd")
    assert r["title"] == "수능특강라이트 영어(개정판)" and r["detail"] == "13강"

    rpu = S.list_ebs_units(cfg, "Reading Power 유형편(완성)")
    assert [(x["label"], x["folder"]) for x in rpu] == [("Unit 18", "분석본")]

    stu = S.list_ebs_units(cfg, "2027 수능특강 영어")
    assert {(x["label"], x["folder"]) for x in stu} == {
        ("3강", "분석본/유형편"), ("Test1 17~20번", "분석본/TEST 1")}


# ── 올림포스 기출 (지문 출처 시트 → 모의고사 분석본) ─────────────────────────

OLPO_G1 = "2026 올림포스 전국연합학력평가 기출 영어독해 고1(2025.10)"


def _write_olpo_sheet(path, columns, red=()):
    """실물 시트 모양: 1행 'N강(유형)', A열 학평Analysis/1..16, 표 아래 메모."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    headers = ["13강(빈칸)", "14강(빈칸)", "15강(무관)"]
    for ci, h in enumerate(headers, start=2):
        ws.cell(1, ci, h)
    for ri in range(17):
        ws.cell(ri + 2, 1, "학평Analysis" if ri == 0 else ri)
    for h, codes in columns.items():
        ci = headers.index(h) + 2
        for ri, c in enumerate(codes):
            cell = ws.cell(ri + 2, ci, int(c))
            if c in red:
                cell.font = Font(color="FFFF0000")
    ws.cell(21, 3, "2025버전에서 사라지는 것들")     # A열 빈 메모 행 — 무시
    ws.cell(22, 3, 220333)
    wb.save(path)


@pytest.fixture
def olpo_tree(tmp_path):
    book = tmp_path / "EBS 작업" / OLPO_G1
    nested = book / "13강" / "13강 통합본" / "분석본"
    nested.mkdir(parents=True)
    (nested / "13강_분석본_통합_창현고1.indd의 사본").write_bytes(b"x")
    old = book / "고1 올림포스 기출 지문 출처 2026_14강 입력완료.xlsx"
    new = book / "고1 올림포스 기출 지문출처 2026_15강.xlsx"
    _write_olpo_sheet(old, {"14강(빈칸)": ["250633", "250934"]})
    _write_olpo_sheet(new, {
        "14강(빈칸)": ["250634", "250934"],
        "15강(무관)": ["250935", "250635", "250335", "241035", "240935",
                     "240635", "240335", "231135", "230935", "220335"],
    }, red=("220335",))
    os.utime(old, (1_700_000_000, 1_700_000_000))
    os.utime(new, (1_800_000_000, 1_800_000_000))
    return {"textbook_base": "", "ebs_base": str(tmp_path / "EBS 작업")}


def test_olpo_units_from_source_sheets(olpo_tree):
    by = {u["label"]: u for u in S.list_ebs_units(olpo_tree, OLPO_G1)}
    assert set(by) == {"14강", "15강"}
    u15 = by["15강"]
    assert u15["kind"] == "olpo" and u15["header"] == "15강(무관)"
    names = [f["name"] for f in u15["files"]]
    assert names == ["Analysis"] + ["%d번" % n for n in range(1, 9)]   # 빨간 글자 제외
    f7 = u15["files"][names.index("7번")]
    assert f7["code"] == "231135" and f7["display"] == "고1 23년 11월 35번"
    assert by["14강"]["files"][0]["code"] == "250634"     # 최신 시트가 이긴다


def test_olpo_resolve_and_range(olpo_tree):
    p = S.resolve_olpo_passage(olpo_tree, OLPO_G1, "15강", "8번")
    assert p == {"code": "230935", "year": "23", "month": "09", "qNum": "35",
                 "grade": "1", "unit_label": "15강"}
    with pytest.raises(ValueError):
        S.resolve_olpo_passage(olpo_tree, OLPO_G1, "15강", "99번")
    assert S.derive_range_olpo(OLPO_G1, "15강", ["7번", "8번"]) == {
        "title": "올림포스 기출 영어독해 고1", "detail": "15강 7번, 8번",
        "abbrev": "올포기출15강"}
    assert S.list_olpo_units(olpo_tree, "올림포스 독해기본2") is None


def test_parse_olpo_passages_use_mock_exam_analysis(engine, olpo_tree,
                                                    tmp_path, monkeypatch):
    from engines.analysis_merge_v2 import engine as eng_mod
    gbase = tmp_path / "모의고사 작업"
    for ym in ("23년 11월", "23년 9월"):
        d = gbase / "고1" / "고1 23년(3,6,9,10,11월)" / f"고1 {ym}"
        d.mkdir(parents=True)
        (d / f"고1 {ym} 분석본.indd").write_bytes(b"x")
    tpl = tmp_path / "tpl.indd"
    tpl.write_bytes(b"t")
    cfg = dict(olpo_tree, gdrive_base=str(gbase), template_path=str(tpl),
               image_dirs=[], file_patterns=["분석본"])
    monkeypatch.setattr(eng_mod, "_load_config", lambda: cfg)
    job = tmp_path / "job"
    job.mkdir()
    (job / "extra_sources.txt").write_text(json.dumps(
        [{"type": "ebs", "book": OLPO_G1, "unit_folder": "15강",
          "file_name": n} for n in ("7번", "8번")], ensure_ascii=False),
        encoding="utf-8")
    meta = json.load(open(engine.parse({}, str(job), lambda m: None),
                          encoding="utf-8"))
    assert [s["code"] for s in meta["sources"]] == ["231135", "230935"]
    assert all(os.path.isfile(c) for s in meta["sources"]
               for c in s["candidates"])
    assert meta["ranges"] == [{"title": "올림포스 기출 영어독해 고1",
                               "detail": "15강 7번, 8번",
                               "abbrev": "올포기출15강"}]


# ── 축약/유도 규칙 (실물 전례 + 사용자 예시로 고정) ──────────────────────────

def test_publisher_short():
    assert S.publisher_short("공통영어1", "공통영어1 비상(홍)") == "비상홍"
    assert S.publisher_short("영어Ⅰ", "영어Ⅰ YBM(박)") == "YBM박"
    assert S.publisher_short("공통영어1", "능률(오)") == "능률오"


def test_join_abbrev_digit_boundary():
    assert S.join_abbrev("비상홍", "3과") == "비상홍3과"
    assert S.join_abbrev("올독기2", "9강") == "올독기2_9강"   # 숫자 인접 → _
    assert S.join_abbrev("리딩파워기본", "13강") == "리딩파워기본13강"


def test_derive_range_textbook():
    r = S.derive_range_textbook("공통영어1", "공통영어1 비상(홍)",
                                "공통영어1 비상(홍) 3과", "3과")
    # title = 과 폴더명 (실물 전례: "공통영어1 능률(오) 2과"), detail은 JSX가
    # 소스 간지 cs '챕터'(과 제목)로 보강 — auto 플래그 필수
    assert r == {"title": "공통영어1 비상(홍) 3과", "detail": "",
                 "abbrev": "비상홍3과", "auto": True}


def test_derive_range_ebs_variants():
    r = S.derive_range_ebs("올림포스 독해기본2", "9강",
                           "올림포스 영어독해 기본2 9강.indd")
    # title = 파일명 표기 (실물 전례: "올림포스 영어독해 기본1")
    assert r["title"] == "올림포스 영어독해 기본2"
    assert r["detail"] == "9강" and r["abbrev"] == "올독기2_9강"
    assert r["auto"] is True

    r2 = S.derive_range_ebs("Reading Power 유형편(기본)", "Unit 13",
                            "Reading Power 유형편(기본) Unit 13 Ex, 1, 2.indd")
    assert r2["title"] == "Reading Power 유형편(기본)"
    assert r2["detail"] == "Unit 13 Ex, 1, 2"
    assert r2["abbrev"] == "리딩파워기본13강"

    r3 = S.derive_range_ebs("2026 수능특강 영어", "12-15강",
                            "2026 수능특강 영어(12-15강) 명여버전(완료).indd",
                            "명여버전 완료")
    assert r3["title"] == "2026 수능특강 영어"
    assert r3["abbrev"] == "26수특12-15강"


def test_default_result_name_mixed_user_example(engine):
    """사용자 예시 고정: 세종1_분석본02(비상홍3과,올독기2_9강)_내지."""
    ranges = [
        S.derive_range_textbook("공통영어1", "공통영어1 비상(홍)",
                                "공통영어1 비상(홍) 3과", "3과"),
        S.derive_range_ebs("올림포스 독해기본2", "9강",
                           "올림포스 영어독해 기본2 9강.indd"),
    ]
    assert engine.default_result_name("세종1", "2", ranges) == \
        "세종1_분석본02(비상홍3과,올독기2_9강)_내지"
    # 모의고사 축약과 혼합
    mixed = [{"title": "고1 25년 3월 모의고사", "detail": "20번"}] + ranges
    assert engine.default_result_name("배재1", "1", mixed) == \
        "배재1_분석본01(25년3월,비상홍3과,올독기2_9강)_내지"


# ── extra_sources 폼 입력 파싱 ───────────────────────────────────────────────

def test_parse_extra_sources():
    raw = json.dumps([
        {"type": "textbook", "subject": "공통영어1",
         "publisher": "공통영어1 비상(홍)",
         "lesson_folder": "공통영어1 비상(홍) 3과",
         "file_name": "공통영어1 비상(홍) 3과.indd"},
        {"type": "ebs", "book": "올림포스 독해기본2",
         "unit_folder": "올림포스 독해기본2 9강",
         "file_name": "올림포스 영어독해 기본2 9강.indd"},
        {"type": "ebs", "book": "2026 수능특강 영어", "unit_folder": "",
         "file_name": "2026 수능특강 영어(Test 2) 함원장버전.indd"},
    ], ensure_ascii=False)
    out = S.parse_extra_sources(raw)
    assert [o["type"] for o in out] == ["textbook", "ebs", "ebs"]
    assert out[2]["unit_folder"] == ""            # 교재 직속(수능특강) 허용
    assert S.parse_extra_sources("") == []
    with pytest.raises(ValueError):
        S.parse_extra_sources("[{broken")
    with pytest.raises(ValueError):
        S.parse_extra_sources('[{"type":"webtoon"}]')
    with pytest.raises(ValueError):                # 필수 필드 누락
        S.parse_extra_sources('[{"type":"textbook","subject":"공통영어1"}]')


def test_parse_requires_codes_or_sources(engine, tmp_path):
    job = tmp_path / "job"
    job.mkdir()
    with pytest.raises(ValueError):
        engine.parse({}, str(job), lambda m: None)


# ── verify doc 게이트 (W28-확장) ─────────────────────────────────────────────

GOOD_REPORT = {
    "unresolved": [],
    "geometry": {"max_drift_before_mm": 4.0, "max_residual_mm": 0.0,
                 "count_mismatch_pages": [], "pages": []},
    "links": {"total": 4, "missing_before": 0, "relinked": 0, "embedded": 0,
              "missing_after": 0, "unembedded_after": 0, "detail": []},
    "ganji": {"ok": True},
    "overflow_pages": [],
    "doc_sources": [
        {"label": "공통영어1 비상(홍) 3과", "type": "textbook",
         "pages_total": 15, "pages_merged": 14, "skipped_pages": [0],
         "page_w": 210.0, "page_h": 297.0, "tpl_w": 210.0, "tpl_h": 297.0,
         "size_mismatch": False},
    ],
}


def _make_job(tmp_path, report, doc_sources_in_settings=1):
    job = tmp_path / "job"
    out = job / "out"
    out.mkdir(parents=True)
    indd = out / "검증_분석본01(비상홍3과)_내지.indd"
    indd.write_bytes(b"INDD" * 64)
    gbase = tmp_path / "gdrive"
    gbase.mkdir()
    settings = {
        "output_indd": str(indd),
        "output_dir": str(out),
        "report_path": str(job / "v2_report.json"),
        "log_path": str(job / "merge_log.txt"),
        "sources": [{"kind": "doc", "label": f"doc{i}"}
                    for i in range(doc_sources_in_settings)],
        "gdrive_save": {"base": str(gbase), "exam_period": "테스트기간",
                        "school_dir": "검증1"},
    }
    (job / ".settings.json").write_text(
        json.dumps(settings, ensure_ascii=False), encoding="utf-8")
    if report is not None:
        (job / "v2_report.json").write_text(
            json.dumps(report, ensure_ascii=False), encoding="utf-8")
    return str(job), str(gbase)


def test_verify_doc_gate_pass(engine, tmp_path):
    job, gbase = _make_job(tmp_path, GOOD_REPORT)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["doc_ok"] is True and res["ok"] is True
    assert res["doc_sources"][0]["pages_merged"] == 14


def test_verify_doc_gate_size_mismatch_blocks(engine, tmp_path):
    bad = json.loads(json.dumps(GOOD_REPORT))
    bad["doc_sources"][0].update(
        {"size_mismatch": True, "page_w": 182.0, "page_h": 257.0,
         "pages_merged": 0})
    job, gbase = _make_job(tmp_path, bad)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["doc_ok"] is False and res["ok"] is False
    assert res["gdrive_path"] is None and not os.listdir(gbase)
    assert any("비호환" in w for w in res["warnings"])


def test_verify_doc_gate_zero_pages_blocks(engine, tmp_path):
    bad = json.loads(json.dumps(GOOD_REPORT))
    bad["doc_sources"][0]["pages_merged"] = 0
    job, _ = _make_job(tmp_path, bad)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["doc_ok"] is False and res["ok"] is False


def test_verify_doc_gate_missing_report_blocks(engine, tmp_path):
    """settings에 doc 소스가 있는데 리포트가 비면 — JSX 중도 실패 봉합."""
    bad = json.loads(json.dumps(GOOD_REPORT))
    bad["doc_sources"] = []
    job, _ = _make_job(tmp_path, bad)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["doc_ok"] is False and res["ok"] is False
    assert any("리포트 누락" in w for w in res["warnings"])


def test_verify_preexisting_overflow_warns_but_passes(engine, tmp_path):
    """소스가 결합 전부터 갖고 있던 오버셋(실측: 올독기2 9강 p13)은 경고로만
    노출 — 게이트 통과·GDrive 저장 진행. 결합이 새로 만든 오버셋은 여전히
    게이트 실패."""
    rep = json.loads(json.dumps(GOOD_REPORT))
    rep["preexisting_overflow_pages"] = ["27"]
    job, gbase = _make_job(tmp_path, rep)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["overflow_ok"] is True and res["ok"] is True
    assert res["preexisting_overflow"] == ["27"]
    assert any("소스 자체 오버셋" in w for w in res["warnings"])
    assert os.listdir(gbase)                      # 저장 진행

    bad = json.loads(json.dumps(GOOD_REPORT))
    bad["overflow_pages"] = ["5"]                 # 결합이 새로 만든 오버셋
    job2, gbase2 = _make_job(tmp_path / "b", bad)
    res2 = engine.verify(job2, lambda m: None)
    assert res2["gates"]["overflow_ok"] is False and res2["ok"] is False


def test_verify_doc_gate_true_when_no_doc_sources(engine, tmp_path):
    good = json.loads(json.dumps(GOOD_REPORT))
    good["doc_sources"] = []
    job, _ = _make_job(tmp_path, good, doc_sources_in_settings=0)
    res = engine.verify(job, lambda m: None)
    assert res["gates"]["doc_ok"] is True and res["ok"] is True


# ── JSX / 폼 계약 ────────────────────────────────────────────────────────────

def test_jsx_doc_source_contract():
    jsx = os.path.join(REPO, "engines", "analysis_merge_v2", "jsx", "run.jsx")
    src = open(jsx, encoding="utf-8").read()
    assert 'srcMeta.kind === "doc"' in src        # doc 소스 분기
    assert "doc_sources" in src                   # 리포트 계약
    assert "size_mismatch" in src                 # 페이지 크기 대조
    assert "enrichDocRange" in src                # 범위 보강 (실물 표기)
    assert '"카테고리"' in src and '"부제목"' in src   # EBS 개요 cs (실측)
    assert '"범위"' in src and '"챕터"' in src         # 교과서 간지 cs (실측)
    assert "skipped_pages" in src                 # 간지/개요 페이지 스킵 기록


def test_template_extra_sources_contract():
    tpl = os.path.join(REPO, "templates", "typeset_analysis_merge_v2.html")
    src = open(tpl, encoding="utf-8").read()
    assert 'name="extra_sources"' in src
    assert "collectExtraSources" in src
    assert "deriveTbRange" in src and "deriveEbsRange" in src
    assert "joinAbbrev" in src                    # 축약 규칙 미러
    assert "/api/analysis-merge/tb-subjects" in src
    assert "/api/analysis-merge/ebs-books" in src


def test_routes_registered():
    """새 폴더 스캔 API 5종이 url_map에 존재 (파리티 픽스처는 별도 재생성)."""
    os.environ.setdefault("YONE_SECRET_KEY", "test-secret-key")
    import app as app_module
    rules = {r.rule for r in app_module.app.url_map.iter_rules()}
    for rule in ("/api/analysis-merge/tb-subjects",
                 "/api/analysis-merge/tb-publishers",
                 "/api/analysis-merge/tb-lessons",
                 "/api/analysis-merge/ebs-books",
                 "/api/analysis-merge/ebs-units"):
        assert rule in rules
