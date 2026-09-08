# -*- coding: utf-8 -*-
"""포켓북 **어휘 파트** — 단어장 원고(xlsx) → 지면 항목.

학교에 따라 포켓북 뒤에 어휘(또는 듣기)가 붙는다. 광문1 은 `어휘끝 고교기본`
Unit 15~23 이 들어간다(26-1-기말 완성본에도 같은 자리에 있었다).

## 원고 모양 (광문보카 xlsx 실측)

시트 하나가 Unit 하나이고, 낱말 하나가 **빈 줄로 끊긴 덩이**다:

    B        C                          D
    DAY15    회독
    1        sensible                   [형] 분별력 있는, 현명한 (↔ senseless …)
             [ˈsɛnsəbəl]
             sensitivity [명] 민감함 …   It wasn’t [sensible] to go out alone …   ← 파생어는 C, 예문은 D
                                        너무 늦은 저녁에 혼자 밖에 나간 것은 …
    (빈 줄)

  · C 칸은 자리에 따라 **표제어 → 발음기호 → 파생어** 로 뜻이 바뀐다.
  · D 칸의 예문은 **영어·한글이 번갈아** 온다. 한 낱말에 예문이 둘일 수 있다.
  · `[…]` 는 그 낱말이 쓰인 자리다 — 지면에서 색이 붙는다.

## 지면 규격 (26-1-기말 광문1 완성본 실측)

    한 면 15낱말 · 한 DAY(30낱말) = 펼침면 하나
    상자 115×177mm · 표제어 10pt / 뜻 7pt(행간 12) / 예문·예문뜻 6pt(행간 7)
    강조색 CMYK 10/85/57/0 = #DB4D5D
"""
import os
import re

_HAN = re.compile(r"[가-힣]")
_IPA = re.compile(r"^\[[^]]*\]$")
_DAY = re.compile(r"^(DAY|Day|Unit)\s*(\d+)", re.I)


def read(path, log=lambda m: None):
    """단어장 xlsx → [{day, no, word, ipa, mean, derived, ex:[(en,ko)…]}].

    시트 차례가 곧 책 차례다(Unit 15 → 16 → …). 시트 이름이 아니라 **B 칸의
    `DAY15`** 를 머리말로 쓴다 — 완성본 꼬리말이 그 말을 쓴다.
    """
    try:
        import openpyxl
    except ImportError:
        log("[어휘] openpyxl 이 없어 어휘 파트를 건너뛴다")
        return []
    if not path or not os.path.exists(path):
        log("[어휘] 원고를 못 찾았다: %s" % path)
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        out += _sheet(rows, name)
    log("■ 어휘 %d낱말 (%d일치)" % (out and len(out) or 0,
                                len({e["day"] for e in out})))
    return out


def _cell(r, i):
    v = r[i] if r and len(r) > i else None
    return "" if v is None else str(v).strip()


def _sheet(rows, sheet):
    out, cur, day = [], None, sheet
    for r in rows:
        b, c, d = _cell(r, 1), _cell(r, 2), _cell(r, 3)
        if _DAY.match(b):                      # `DAY15 / 회독` — 덩이 머리
            day = b
            continue
        if not (b or c or d):                  # 빈 줄 — 덩이 끝
            if cur:
                out.append(cur)
                cur = None
            continue
        if b.isdigit():                        # 번호 줄 = 표제어와 뜻
            if cur:
                out.append(cur)
            cur = {"day": day, "no": int(b), "word": c, "ipa": "",
                   "mean": d, "derived": "", "ex": []}
            continue
        if cur is None:
            continue
        if c and not cur["ipa"] and _IPA.match(c):
            cur["ipa"] = c
        elif c:
            cur["derived"] = (cur["derived"] + " " + c).strip()
        if d:
            # 예문은 영어 → 한글 차례로 온다. 한글이면 앞 예문의 짝이다.
            if cur["ex"] and cur["ex"][-1][1] == "" and _HAN.search(d):
                cur["ex"][-1] = (cur["ex"][-1][0], d)
            else:
                cur["ex"].append((d, ""))
    if cur:
        out.append(cur)
    return out


# 낱말 하나가 차지하는 높이(mm) 어림 — **CSS 행간 그대로**여야 한다.
# 처음엔 눈대중으로 잡았다가 35면이 전부 판면 밖으로 15mm 씩 넘쳤다.
_W_LEAD = 10.0 * 1.18 / 72.0 * 25.4     # 표제어 10pt · line-height 1.18
_M_LEAD = 12.0 / 72.0 * 25.4            # 뜻 행간 12pt
_X_LEAD = 7.0 / 72.0 * 25.4             # 예문·예문뜻 행간 7pt
_GAP = 1.4                              # 낱말 사이(.e margin)


# ⚠ 줄 수는 **실제 자폭으로** 재야 한다. 한글 1em·영문 0.5em 으로 어림했더니
#   같은 면이 어떤 곳은 1.12배, 어떤 곳은 1.52배로 어긋나 35면이 전부 판면
#   밖으로 밀렸다. 지면 글꼴(Noto Sans CJK KR)의 자폭을 그대로 읽는다.
_FONT = "NotoSansCJKkr-Medium.otf"
_ADV = {}


def _load_adv():
    import glob
    try:
        from fontTools.ttLib import TTFont
        here = os.path.dirname(os.path.abspath(__file__))
        # …/core/render_vivlio/series/pocketbook → 저장소 뿌리까지 네 단계다.
        root = here
        for _ in range(4):
            root = os.path.dirname(root)
        hit = glob.glob(os.path.join(root, "**", _FONT), recursive=True)
        f = TTFont(hit[0])
        em, hm, cm = f["head"].unitsPerEm, f["hmtx"], f.getBestCmap()
        for code, g in cm.items():
            _ADV[chr(code)] = hm[g][0] / em
    except Exception:
        pass
    _ADV["_"] = None


def _lines(txt, width_mm, size_pt, track_em=-0.025):
    """이 글이 몇 줄이 될지 — 글꼴 자폭으로 잰다."""
    if not txt:
        return 0
    if "_" not in _ADV:
        _load_adv()
    em = size_pt / 72.0 * 25.4
    w = 0.0
    for ch in txt:
        a = _ADV.get(ch)
        if a is None:
            a = 1.0 if ord(ch) > 0x2000 else 0.55
        w += (a + track_em) * em
    return max(1, int(w // width_mm) + (1 if w % width_mm else 0))


def height_mm(e, width_mm=115.0):
    """낱말 하나의 높이(mm). **실측값(`_h`)이 있으면 그것이 이긴다.**"""
    if e.get("_h"):
        return e["_h"]
    h = _W_LEAD + _lines(e.get("mean", ""), width_mm, 7.0) * _M_LEAD
    if e.get("derived"):
        h += _lines(e["derived"], width_mm, 6.0) * _X_LEAD
    for en, ko in e.get("ex") or ():
        h += _lines(en, width_mm, 6.0) * _X_LEAD
        if ko:
            h += _lines(ko, width_mm, 6.0) * _X_LEAD
    return h + _GAP


def label(book, day):
    """머리말 — 완성본이 쓰는 말(`어휘 고교기본 Unit 6`)."""
    import re as _re
    m = _re.search(r"(\d+)", day or "")
    return "어휘 %s Unit %s" % (book, m.group(1) if m else day)


def number(entries):
    """번호는 **어휘 파트 전체에 이어지는 일련번호**다(001, 002 …)."""
    for i, e in enumerate(entries, 1):
        e["seq"] = i
    return entries


def pages(entries, box_h=175.0, width_mm=114.8):
    """낱말 목록 → 면 목록.

    ⚠ **면당 개수를 정해 놓으면 안 된다.** 낱말마다 뜻 길이도, 예문 수도(하나
      또는 둘), 파생어 유무도 달라서 같은 15개라도 높이가 40mm 넘게 차이 난다.
      개수로 끊었더니 아래쪽 낱말이 판면 밖으로 밀려 나갔다. **높이로 채운다.**

    ⚠ **DAY 가 바뀌면 면을 새로 연다.** 이어 붙이면 오늘 몫이 어디까지인지
      학생이 알 수 없다(완성본도 DAY 단위로 면을 끊는다).
    """
    out = []
    for day, group in _by_day(entries):
        hs = [height_mm(e, width_mm) for e in group]
        tot = sum(hs)
        # 어림에는 오차가 있다 — 딱 맞게 잡으면 마지막에 한 낱말이 넘쳐
        # **낱말 하나짜리 면**이 생긴다. 92% 로 잡아 여유를 둔다.
        room = box_h * 0.92
        n = max(1, int(tot // room) + (1 if tot % room else 0))
        # ⚠ 꽉 채우고 남은 것을 뒷면에 밀면 **낱말 한두 개짜리 면**이 생긴다.
        #   DAY 가 쓸 면 수를 먼저 정하고 그만큼으로 **고르게** 나눈다.
        target = tot / n
        cur, used, made = [], 0.0, 0
        for e, h in zip(group, hs):
            over = used + h > box_h                    # 상자를 넘는다
            fair = used >= target - h / 2 and made < n - 1   # 제 몫을 채웠다
            if cur and (over or fair):
                out.append({"day": day, "items": cur})
                made += 1
                cur, used = [], 0.0
            cur.append(e)
            used += h
        if cur:
            out.append({"day": day, "items": cur})
    return out


def _by_day(entries):
    day, group = None, []
    for e in entries:
        if day is not None and e["day"] != day:
            yield day, group
            group = []
        day = e["day"]
        group.append(e)
    if group:
        yield day, group
