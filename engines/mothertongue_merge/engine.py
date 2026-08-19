"""마더텅 결합 엔진.

엑셀에 정의된 강별 지문 출처 리스트를 보고, 해당 코드들을 모아 분석본
마스터 템플릿 위에 결합한다.

소스 우선순위:
  1) 모의고사 작업/고{N}/고{N} {YY}년 {M}월/*분석본*.indd
     (analysis_merge 엔진과 동일 — 마커 텍스트 매칭)
  2) Fallback: {fallback_base}/{N}강 .../*.indd 의 고정 위치
     (엑셀 행 순서 → 페이지 인덱스 = 1 + 2 * (rowOffset) 0-기반)

엑셀 형식 (header_row=3, first_data_row=4 기준):
  - 헤더 행: 첫 셀은 None, 나머지는 "{N:02d} {강명}" 형식
  - 데이터 행: A 열 = 강 내 순번, B~ 열 = 6자리 YYMMQQ 코드 (또는 빈셀)
"""

import json
import os
import re
import unicodedata

from config import STUDIO_ROOT
from core.base import BaseEngine, safe_copy_indd


def _nfc(s):
    return unicodedata.normalize("NFC", s)


_DIR = os.path.dirname(os.path.abspath(__file__))
_CONFIG_PATH = os.path.join(_DIR, "config.json")
_PATH_KEYS = ["template_path", "gdrive_base", "output_base"]


def _load_config():
    from core.paths import resolve_config
    with open(_CONFIG_PATH, encoding="utf-8") as f:
        cfg = json.load(f)
    return resolve_config(cfg, _PATH_KEYS)


def _expand_path(p):
    if not p:
        return p
    return os.path.expanduser(p)


_LESSON_HEADER_RE = re.compile(r"^\s*(\d{1,2})\s*(.*)$")


def _parse_excel_lessons(excel_path, header_row=3, first_data_row=4):
    """엑셀에서 강 목록과 강별 코드 리스트 반환.

    Returns:
        list of dict: [
            {
              "lesson_num": "03",
              "lesson_name": "필자의 주장",
              "column": 2,          # openpyxl 1-기반 column index
              "header_label": "03 필자의 주장",
              "passages": [
                  {"row_offset": 0, "code": "221120"},
                  {"row_offset": 1, "code": "230920"},
                  ...
              ]
            },
            ...
        ]
    """
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    lessons = []
    # 헤더 행 스캔 — 2열부터
    max_col = ws.max_column
    for col in range(2, max_col + 1):
        cell = ws.cell(header_row, col).value
        if cell is None:
            continue
        text = str(cell).strip()
        m = _LESSON_HEADER_RE.match(text)
        if not m:
            continue
        lesson_num = m.group(1).zfill(2)  # "3" → "03"
        lesson_name = m.group(2).strip()
        # 코드 수집 (first_data_row 부터, 빈셀이면 skip)
        passages = []
        row_offset = 0
        for r in range(first_data_row, ws.max_row + 1):
            v = ws.cell(r, col).value
            if v is None or str(v).strip() == "":
                row_offset += 1
                continue
            code_str = str(v).strip()
            # 숫자값이면 zfill, 문자열이면 그대로
            if isinstance(v, (int, float)):
                code_str = str(int(v)).zfill(6)
            if not re.match(r"^\d{6}$", code_str):
                row_offset += 1
                continue
            passages.append({"row_offset": row_offset, "code": code_str})
            row_offset += 1
        lessons.append({
            "lesson_num": lesson_num,
            "lesson_name": lesson_name,
            "column": col,
            "header_label": text,
            "passages": passages,
        })
    return lessons


def _gdrive_subdir(base, name_nfc):
    """GDrive 원본 폴더(base) 안에서 name_nfc 와 같은 이름의 하위 폴더 경로.

    디스크 이름이 NFD 라 단순 join 이 빗나갈 수 있어 NFC 비교로 되찾는다.
    """
    if not base or not os.path.isdir(base):
        return None
    direct = os.path.join(base, name_nfc)
    if os.path.isdir(direct):
        return direct
    try:
        for raw in os.listdir(base):
            if _nfc(raw) == name_nfc:
                cand = os.path.join(base, raw)
                if os.path.isdir(cand):
                    return cand
    except OSError:
        pass
    return None


# 파일명이 특정 지문 하나를 가리키는지 — JSX 마커와 같은 식별자
# ("고1 16년 6월 36번 분석본.indd"). 이 이름은 후보를 고르는 힌트일 뿐,
# 실제 페이지 선택은 JSX 의 마커 매칭이 결정한다.
_PASSAGE_NAME_RE = re.compile(
    r"(?:고\d+\s*)?\d{1,2}\s*년\s*\d{1,2}\s*월\s*\d{1,2}\s*(?:[-~]\s*\d{1,2})?\s*번"
)


def _lesson_indd_entries(fallback_base, lesson_num):
    """마더텅 {N}강 폴더 최상위의 .indd 목록 → [(name_nfc, full_path), ...].

    잠금/백업/템플릿 사본은 제외. 강 파일과 지문별 파일을 모두 포함한다
    (분류는 호출부에서 `_PASSAGE_NAME_RE` 로).
    """
    if not fallback_base:
        return []
    # 미러 ∪ GDrive 합집합 — 미러가 stale해도 강 폴더 누락 방지
    from engines.final.extractor import list_manuscript_entries
    lesson_int = int(lesson_num)
    # 폴더명: "{N}강 {lesson_name}" — N 은 0-padded 가 아니므로 lesson_int 와 매칭
    target_prefixes = [f"{lesson_int}강", f"{lesson_int:02d}강"]
    entries = list_manuscript_entries(fallback_base, None)
    matched = [(nfc, full) for nfc, full in entries
               if any(nfc.startswith(p) for p in target_prefixes)]
    if not matched:
        # 2차 (배치 F): 숫자 토큰 매칭 — 접두가 아니어도 'N강' 토큰이
        # 있으면 인정 (예: '마더텅 7강 …', '제7강'). "1강"↛"11강" 오매칭
        # 방지를 위해 숫자 경계 가드. 1차 0건일 때만 진입 (기존 무변경).
        _pat = re.compile(rf"(?<!\d)0*{lesson_int}\s*강(?!\d)")
        matched = [(nfc, full) for nfc, full in entries
                   if _pat.search(nfc)]
    out = []
    for nfc, full in matched:
        if not os.path.isdir(full):
            continue
        # 강 폴더 **안쪽**도 미러 ∪ GDrive 합집합으로 본다. 종전에는
        # os.listdir(full) 이었는데 full 이 미러 경로라, 미러가 아직 못 받은
        # 새 파일이 통째로 안 보였다(2026-08-18 13강 지문별 분석본 3개 실측:
        # 미러 0개 / GDrive 3개).
        sub = _gdrive_subdir(fallback_base, nfc) or full
        for nfc_f, full_f in list_manuscript_entries(sub, None):
            if not nfc_f.lower().endswith(".indd"):
                continue
            if nfc_f.startswith("~") or "idlk" in nfc_f.lower():
                continue
            # "분석본_내지(마스터).indd" 같은 템플릿 사본 제외
            if "마스터" in nfc_f or "템플릿" in nfc_f:
                continue
            out.append((nfc_f, full_f))
    out.sort(key=lambda t: t[0])
    return out


def _find_lesson_passage_files(fallback_base, lesson_num):
    """강 폴더 안의 **지문별 분석본** .indd (예: "고1 16년 6월 36번 분석본.indd").

    마더텅 옛날 지문은 월 폴더(모의고사 작업/고1/고1 YY년 M월)가 아예 없어서
    1차 후보가 0개다. 이 파일들을 primary 후보로 붙여 JSX 마커 매칭이
    페이지를 고르게 한다 — 실측 구조는 월 폴더 분석본과 동일
    (p1 간지-마스터 / p2 마커 / p3 이어지는 면).
    """
    return [full for nfc, full in _lesson_indd_entries(fallback_base, lesson_num)
            if _PASSAGE_NAME_RE.search(nfc)]


def _find_fallback_file(fallback_base, lesson_num):
    """강 폴더의 **강 단위** 파일 하나 (고정 페이지 인덱스용).

    지문 하나를 가리키는 파일(_PASSAGE_NAME_RE)은 제외한다 — 강 파일로
    잘못 잡히면 row_offset 기반 고정 인덱스가 엉뚱한 페이지를 가리켜
    조용히 틀린 지문이 실린다(2026-08-18 13강에 지문별 파일이 들어오며 실제
    위험 발생). 지문별 파일은 `_find_lesson_passage_files` 로 primary 후보에
    붙어 마커로 판정된다.
    """
    for nfc, full in _lesson_indd_entries(fallback_base, lesson_num):
        if _PASSAGE_NAME_RE.search(nfc):
            continue
        return full
    return None


def _find_month_folder(gdrive_base, grade_num, year, month):
    """월 폴더 매칭 — analysis_merge와 같은 눈 (배치 F 합류).

    종전엔 analysis_merge의 1차(리터럴)만 사본으로 갖고 있어 배치 A(W33)의
    2차 숫자 토큰 매칭(언더스코어/4자리 연도 — 감사 실증 #6)을 못 받았다.
    위임으로 단일화 — 1차 로직은 동일했으므로 기존 매칭분 무변경.
    """
    from engines.analysis_merge.engine import AnalysisMergeEngine
    return AnalysisMergeEngine._find_month_folder(
        gdrive_base, grade_num, year, month)


def _find_analysis_files(month_folder, patterns=None, log=None):
    """분석본 후보 매칭 — analysis_merge와 같은 눈 (배치 F 합류).

    1차(리터럴 substring)는 동일, 2차('분석.indd' 축약형 등 정규화 폴백 —
    배치 A 미해결 (4))를 위임으로 함께 받는다.

    3차 — 구조 폴백: **파일명은 우선순위 힌트일 뿐 판정 기준이 아니다.**
    월 폴더 최상위의 나머지 .indd 도 뒤에 붙여 후보로 넘기고, 실제 채택은
    JSX 의 마커("고N YY년 M월 QQ번") 매칭이 결정한다. 23년 폴더가
    '분석본' 대신 '고1 23년 6월 모의고사(New).indd' 로 명명돼 있어 8개 코드가
    통째로 '해결 불가' 로 빠진 사례(2026-08-18) 방지 — 실측 결과 이 파일들은
    24/25년 분석본과 동일 구조(37p, 짝수 페이지마다 마커)였다.
    이름 매칭분이 항상 앞이라 기존 매칭 순서는 바뀌지 않는다.
    """
    from engines.analysis_merge.engine import AnalysisMergeEngine
    named = AnalysisMergeEngine._find_analysis_files(month_folder, patterns)
    if not month_folder:
        return named

    from engines.final.extractor import list_manuscript_entries
    seen = set(named)
    extras = []
    for nfc, full in list_manuscript_entries(month_folder, None):
        if full in seen or not nfc.lower().endswith(".indd"):
            continue
        if nfc.startswith("~") or "idlk" in nfc.lower():
            continue
        # 템플릿/마스터 사본은 분석본이 아니다
        if "마스터" in nfc or "템플릿" in nfc:
            continue
        extras.append(full)
    extras.sort(key=lambda p: (len(os.path.basename(p)), os.path.basename(p)))
    if extras and log:
        # 침묵 금지 — 이름이 안 맞아 들어온 후보를 항상 노출
        log("    [후보 보강] 이름 비매칭 .indd " + str(len(extras)) + "개 추가: "
            + ", ".join(os.path.basename(e) for e in extras[:4]))
    return named + extras


class MothertongueMergeEngine(BaseEngine):

    @property
    def material_id(self):
        return "mothertongue_merge"

    @property
    def display_name(self):
        return "마더텅 결합"

    @property
    def description(self):
        return "마더텅 지문 출처 엑셀에 따라 강 단위로 분석본을 결합 → .indd + PDF"

    @property
    def subject_prefix(self):
        return "마더텅결합"

    @property
    def template_name(self):
        return "typeset_mothertongue_merge.html"

    @property
    def log_filename(self):
        return "merge_log.txt"

    @property
    def jsx_timeout(self):
        return 1200  # 20분 — 여러 .indd 열기 + 큰 강 (11강 72지문) 대비

    @property
    def completion_marker(self):
        """JSX 완료 표식 — 러너가 이 줄을 보면 즉시 끝낸다.

        선언하지 않으면 러너는 로그에 "PDF:"가 뜨는 것으로 완료를 판정한다.
        PDF 추출을 끄면(export_pdf=false) 그 줄이 영원히 안 나와 정상 완료가
        **스톨(300초 미갱신)로 오보고**된다 — 2026-08-19 웹 실행 실측
        (JSX는 12:21:02에 정상 종료했는데 잡은 '작업 보류'로 실패 처리).
        """
        return "===== 완료 ====="

    @property
    def upload_spec(self):
        return {
            "slots": [
                {
                    "key": "lessons",
                    "label": "선택한 강 (콤마 구분)",
                    "type": "text",
                    "required": True,
                },
                {
                    "key": "issue_num",
                    "label": "분석본 순번(주차) — 간지·파일명에 함께 쓰임",
                    "type": "text",
                    "required": False,
                },
                {
                    "key": "result_name",
                    "label": "결과 파일명 (확장자 제외 — 비우면 자동 생성)",
                    "type": "text",
                    "required": False,
                },
            ]
        }

    # ─── 헬퍼 ───────────────────────────────────────────────────────

    def _get_meta(self, job_dir, key, default=""):
        path = os.path.join(job_dir, f"{key}.txt")
        if os.path.exists(path):
            return open(path, encoding="utf-8").read().strip()
        return default

    @staticmethod
    def list_lessons(grade, school):
        """API 용 — 주어진 학년+학교의 강 목록 반환."""
        cfg = _load_config()
        key = f"{grade}-{school}"
        scfg = cfg["schools"].get(key)
        if not scfg:
            return {"lessons": [], "error": f"엑셀 설정 없음: {key}"}
        excel_path = _expand_path(scfg["excel_path"])
        if not os.path.exists(excel_path):
            return {"lessons": [], "error": f"엑셀 파일 없음: {excel_path}"}
        try:
            lessons = _parse_excel_lessons(
                excel_path,
                header_row=scfg.get("header_row", 3),
                first_data_row=scfg.get("first_data_row", 4),
            )
        except Exception as e:
            return {"lessons": [], "error": f"엑셀 파싱 실패: {e}"}
        return {
            "lessons": [
                {
                    "lesson_num": l["lesson_num"],
                    "lesson_name": l["lesson_name"],
                    "header_label": l["header_label"],
                    "count": len(l["passages"]),
                }
                for l in lessons
            ]
        }

    # ─── parse ─────────────────────────────────────────────────────

    def parse(self, uploaded_files, job_dir, log):
        cfg = _load_config()
        patterns = cfg.get("file_patterns", ["분석본"])

        grade = self._get_meta(job_dir, "grade", "고1")
        school = _nfc(self._get_meta(job_dir, "school", "한영고1"))
        grade_num = grade.replace("고", "")

        school_key = f"{grade}-{school}"
        scfg = cfg["schools"].get(school_key)
        if not scfg:
            raise ValueError(f"엑셀 설정 없음 (config.json schools 에 추가 필요): {school_key}")

        excel_path = _expand_path(scfg["excel_path"])
        fallback_base = _expand_path(scfg.get("fallback_base", ""))
        pages_per_passage = scfg.get("pages_per_passage", 2)

        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"엑셀 파일 없음: {excel_path}")

        # 선택한 강 (콤마 구분 lesson_num 또는 header_label)
        lessons_raw = uploaded_files.get("lessons", "")
        selected = [s.strip() for s in lessons_raw.split(",") if s.strip()]
        if not selected:
            raise ValueError("강이 선택되지 않았습니다.")

        log(f"[엑셀] {os.path.basename(excel_path)}")
        log(f"[선택 강] {', '.join(selected)}")

        all_lessons = _parse_excel_lessons(
            excel_path,
            header_row=scfg.get("header_row", 3),
            first_data_row=scfg.get("first_data_row", 4),
        )

        # 선택된 강만 필터 — lesson_num 또는 header_label 일치
        selected_lessons = []
        for l in all_lessons:
            if (l["lesson_num"] in selected or
                l["lesson_name"] in selected or
                l["header_label"] in selected or
                str(int(l["lesson_num"])) in selected):
                selected_lessons.append(l)

        if not selected_lessons:
            raise ValueError(f"선택한 강을 엑셀에서 찾지 못함: {selected}")

        log(f"[강] {len(selected_lessons)}개 매칭됨")
        for l in selected_lessons:
            log(f"  - {l['header_label']}: {len(l['passages'])}개 지문")

        # 각 강의 코드별로 후보 파일 결정
        gdrive_base = _expand_path(cfg["gdrive_base"])

        # 후보 캐시 — 같은 파일 중복 복사 방지
        local_copied = {}      # gdrive_path → local_path

        # ym 캐시 — 같은 (year, month) 월 폴더는 한 번만 스캔
        ym_cache = {}          # (year, month) → list[local_path]
        ym_folder_cache = {}   # (year, month) → 월 폴더 경로 (후보 리포트용)

        # fallback 캐시 — 강 파일 한 번만 복사
        fallback_cache = {}    # lesson_num → local_path

        # 강 폴더의 지문별 분석본 캐시 — 강마다 한 번만 스캔·복사
        lesson_passage_cache = {}   # lesson_num → list[local_path]

        # GDrive 파일은 EDEADLK 회피를 위해 7-경로 폴백 + 지수 백오프 사용
        # (.claude/rules/gdrive_copy_robust.md 참조)
        from engines.final.extractor import _gdrive_safe_copy

        def _local_copy(src_path):
            if src_path in local_copied:
                return local_copied[src_path]
            base = os.path.basename(src_path)
            local_path = os.path.join(job_dir, f"src_{base}")
            # 같은 이름이 있으면 _2, _3 suffix
            if os.path.isfile(local_path) and os.path.getsize(local_path) > 0:
                b, ext = os.path.splitext(local_path)
                n = 2
                while (os.path.isfile(f"{b}_{n}{ext}") and
                       os.path.getsize(f"{b}_{n}{ext}") > 0):
                    n += 1
                local_path = f"{b}_{n}{ext}"
            ok, err = _gdrive_safe_copy(src_path, local_path, log=log)
            if not ok:
                raise RuntimeError(f"GDrive 파일 복사 실패: {base} — {err}")
            local_copied[src_path] = local_path
            return local_path

        passage_entries = []
        for lesson in selected_lessons:
            ln = lesson["lesson_num"]
            for passage in lesson["passages"]:
                code = passage["code"]
                year = code[0:2]
                month = code[2:4]
                qNum = code[4:6]
                row_offset = passage["row_offset"]

                # 1차 후보: 모의고사 작업/고{N}/고{N} {YY}년 {M}월/*분석본*.indd
                key = (year, month)
                if key not in ym_cache:
                    folder = _find_month_folder(gdrive_base, grade_num, year, month)
                    ym_folder_cache[key] = folder
                    if not folder:
                        ym_cache[key] = []
                    else:
                        files = _find_analysis_files(folder, patterns, log=log)
                        locals_ = []
                        for f in files:
                            locals_.append(_local_copy(f))
                            log(f"    로컬 복사: {os.path.basename(f)}")
                        ym_cache[key] = locals_
                primary_candidates = list(ym_cache[key])

                # 1.5차 후보: 마더텅 {N}강 폴더의 **지문별** 분석본
                #   옛날 지문(16/17/19년 등)은 월 폴더 자체가 없어 1차가 0개다.
                #   월 폴더 후보 뒤에 붙이므로 기존 매칭 순서는 바뀌지 않고,
                #   채택 여부는 JSX 마커 매칭이 결정한다.
                if ln not in lesson_passage_cache:
                    lp_srcs = _find_lesson_passage_files(fallback_base, ln)
                    lp_locals = []
                    for lp in lp_srcs:
                        lp_locals.append(_local_copy(lp))
                        log(f"    지문별 분석본 로컬 복사: {os.path.basename(lp)} (강 {ln})")
                    lesson_passage_cache[ln] = lp_locals
                for lp_local in lesson_passage_cache[ln]:
                    if lp_local not in primary_candidates:
                        primary_candidates.append(lp_local)

                # 2차 후보 (fallback): {fallback_base}/{N}강 .../*.indd 의 고정 위치
                fallback_info = None
                if ln not in fallback_cache:
                    fb_src = _find_fallback_file(fallback_base, ln)
                    if fb_src:
                        fb_local = _local_copy(fb_src)
                        log(f"    fallback 로컬 복사: {os.path.basename(fb_src)} (강 {ln})")
                        fallback_cache[ln] = fb_local
                    else:
                        fallback_cache[ln] = None
                fb_local = fallback_cache[ln]
                if fb_local:
                    # 강 파일 안에서의 위치: 1페이지(간지)+ row_offset 단위 * pages_per_passage
                    # 0-인덱스 페이지: 1 + row_offset * pages_per_passage
                    fixed_start = 1 + row_offset * pages_per_passage
                    fallback_info = {
                        "local_path": fb_local,
                        "start_idx": fixed_start,
                        "page_count": pages_per_passage,
                    }

                if not primary_candidates and not fallback_info:
                    # 침묵 금지 (배치 F): 후보 수준 리포트 동봉
                    _mf = ym_folder_cache.get(key)
                    if _mf:
                        from engines.final.extractor import list_manuscript_entries
                        _names = sorted(
                            n for n, _f in
                            (list_manuscript_entries(_mf, None) or [])
                            if n.lower().endswith(".indd")
                            and not n.startswith("~")
                        )[:8]
                        _extra = (f" (월 폴더 {os.path.basename(_mf)} .indd 후보: "
                                  + (", ".join(_names) if _names else "없음") + ")")
                    else:
                        _extra = " (월 폴더 미발견)"
                    log(f"  [경고] 해결 불가: 강 {ln} 코드 {code} — "
                        f"1차/폴백 모두 없음{_extra}")

                # 헤더(머리말) — "13강 1번 | 고1 23년 9월 36번"
                #   라벨칸=강+마더텅 지문번호(엑셀 순서), 코드칸=원래 출처
                #   지문번호 그대로 유지(2026-08-19 사용자 확정).
                #   header_code 를 비우면 JSX 가 코드칸을 건드리지 않는다.
                try:
                    _ln_int = int(str(ln).lstrip("0") or "0")
                except ValueError:
                    _ln_int = ln
                passage_entries.append({
                    "header_label": f"{_ln_int}강 {row_offset + 1}번",
                    "header_code": "",
                    "code": code,
                    "year": year,
                    "month": month,
                    "qNum": qNum,
                    "lesson_num": ln,
                    "lesson_name": lesson["lesson_name"],
                    "lesson_label": lesson["header_label"],
                    "row_offset": row_offset,
                    "primary_candidates": primary_candidates,
                    "fallback": fallback_info,
                })

        log(f"[지문] 총 {len(passage_entries)}개")

        meta_path = os.path.join(job_dir, "metadata.json")
        with open(meta_path, "w", encoding="utf-8") as f:
            json.dump(
                {"passages": passage_entries, "selected_lessons": selected_lessons},
                f, ensure_ascii=False, indent=2
            )
        return meta_path

    # ─── prepare_jsx ───────────────────────────────────────────────

    @staticmethod
    def _school_abbrev(school):
        """폴더/엑셀용 학교명 → 파일명 축약. "한영고1" → "한영1"."""
        return re.sub(r"고(\d)$", r"\1", _nfc(school or "").strip())

    def default_result_name(self, school, issue_num, lessons_sel):
        """명명 규칙: {학교축약}_분석본{NN}(마더텅{강}강)_내지 — 공백 없음.

        예: 한영1_분석본02(마더텅13강)_내지
        (사용자가 파일명을 직접 타이핑하지 않아도 되게 — 2026-08-18 요청)
        """
        try:
            nn = f"{int(issue_num):02d}"
        except (TypeError, ValueError):
            nn = "01"
        nums = []
        for l in lessons_sel:
            try:
                n = str(int(str(l.get("lesson_num", "")).lstrip("0") or "0"))
            except ValueError:
                n = str(l.get("lesson_num", ""))
            if n and n not in nums:
                nums.append(n)
        body = f"마더텅{','.join(nums)}강" if nums else "마더텅"
        name = f"{self._school_abbrev(school)}_분석본{nn}({body})_내지"
        name = re.sub(r'[/\\:*?"<>|]', "", name)
        return re.sub(r"\s+", "", name)

    def prepare_jsx(self, json_path, job_dir, log):
        cfg = _load_config()
        grade = self._get_meta(job_dir, "grade", "고1")
        school = _nfc(self._get_meta(job_dir, "school", "한영고1"))
        exam_period = self._get_meta(job_dir, "exam_period", "기간미지정")
        issue_num = self._get_meta(job_dir, "issue_num", "").strip()

        with open(json_path, encoding="utf-8") as f:
            _meta_pre = json.load(f)
        lessons_sel = _meta_pre.get("selected_lessons") or []

        # 간지 범위 — "마더텅 13강" / "문장 배열" (2026-08-18 사용자 확정)
        ranges = []
        for l in lessons_sel:
            try:
                n = int(str(l.get("lesson_num", "")).lstrip("0") or "0")
            except ValueError:
                n = l.get("lesson_num", "")
            ranges.append({
                "title": f"마더텅 {n}강",
                "detail": (l.get("lesson_name") or "").strip(),
            })

        # 간지 학교 표기 — "2026 한영고1" (연도는 시험 기간 앞 두 자리)
        yy = (exam_period or "")[:2]
        year_full = f"20{yy}" if yy.isdigit() else ""
        ganji_school = f"{year_full} {school}".strip()

        result_name = self._get_meta(job_dir, "result_name", "").strip()
        if not result_name:
            # 자동 생성 — 학교축약_분석본NN(마더텅N강)_내지
            result_name = self.default_result_name(school, issue_num, lessons_sel)
            log(f"[파일명 자동] {result_name}")

        output_dir = os.path.join(
            STUDIO_ROOT, "조판", "교재", "마더텅결합",
            grade, exam_period, school
        )
        os.makedirs(output_dir, exist_ok=True)
        log(f"[출력 폴더] {output_dir}")

        indd_out = os.path.join(output_dir, f"{result_name}.indd")
        if os.path.exists(indd_out):
            n = 2
            while n < 1000:
                trial = os.path.join(output_dir, f"{result_name}_{n}.indd")
                if not os.path.exists(trial):
                    indd_out = trial
                    break
                n += 1
        pdf_out = os.path.splitext(indd_out)[0] + ".pdf"

        meta = _meta_pre

        settings = {
            "template_path": cfg["template_path"],
            "output_indd": indd_out,
            "output_pdf": pdf_out,
            "output_dir": output_dir,
            "log_path": os.path.join(output_dir, self.log_filename),
            "log_paths": [os.path.join(job_dir, self.log_filename)],
            "keep_head": cfg.get("keep_head", 1),
            # PDF 추출 스위치 — 현재 꺼둠(사용자 요청 2026-08-19).
            # 켜면 도련 포함으로 고정 출력된다(JSX useDocumentBleedWithPDF).
            "export_pdf": bool(cfg.get("export_pdf", True)),
            "ganji_school": ganji_school,
            "issue_num": issue_num or "1",
            "ranges": ranges,
            "school": school,
            "grade": grade,
            "exam_period": exam_period,
            "passages": meta["passages"],
        }
        settings_path = os.path.join(job_dir, ".settings.json")
        with open(settings_path, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)

        jsx_src = os.path.join(_DIR, "jsx", "run.jsx")
        jsx_dst = os.path.join(job_dir, "run.jsx")
        with open(jsx_src, encoding="utf-8") as f:
            jsx_code = f.read()
        jsx_code = jsx_code.replace("__SETTINGS_PATH__", settings_path)
        with open(jsx_dst, "w", encoding="utf-8") as f:
            f.write(jsx_code)

        log(f"[결과 .indd] {os.path.basename(indd_out)}")
        if cfg.get("export_pdf", True):
            log(f"[결과 PDF]  {os.path.basename(pdf_out)}")
        else:
            log("[결과 PDF]  (추출 꺼짐 — .indd 만 생성)")
        return jsx_dst

    # ─── verify ────────────────────────────────────────────────────

    def verify(self, job_dir, log):
        settings_path = os.path.join(job_dir, ".settings.json")
        with open(settings_path, encoding="utf-8") as f:
            settings = json.load(f)

        indd_out = settings["output_indd"]
        pdf_out = settings["output_pdf"]

        export_pdf = bool(settings.get("export_pdf", True))

        warnings = []
        if not os.path.exists(indd_out):
            warnings.append(f".indd 미생성: {os.path.basename(indd_out)}")
        if export_pdf and not os.path.exists(pdf_out):
            warnings.append(f"PDF 미생성: {os.path.basename(pdf_out)}")

        # job 전용 로그(이번 실행분만) 우선 — 공용 로그 누적 오염 방지
        log_paths = settings.get("log_paths") or []
        log_path = log_paths[0] if log_paths else settings.get("log_path", "")
        warn_count = 0
        if log_path and os.path.exists(log_path):
            try:
                content = open(log_path, encoding="utf-8", errors="ignore").read()
                warn_count = len(re.findall(r"\[경고\]|\[오류\]", content))
                if warn_count > 0:
                    warnings.append(f"JSX 경고/오류 {warn_count}건 (로그 확인)")
            except OSError:
                pass

        # PDF 추출이 꺼져 있으면 .indd 생성 여부로 판정한다
        ok = (os.path.exists(pdf_out) if export_pdf
              else os.path.exists(indd_out)) and len(warnings) == 0

        return {
            "ok": ok,
            "overflow": 0,
            "stories": {},
            "pdf_path": pdf_out if os.path.exists(pdf_out) else None,
            "pdf_paths": [pdf_out] if os.path.exists(pdf_out) else [],
            "warnings": warnings,
        }
