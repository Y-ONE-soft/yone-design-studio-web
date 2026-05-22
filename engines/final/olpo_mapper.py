"""고2 올림포스 기출 — YYMMQQ ↔ (강, 본문N) 매핑 + 중요어휘 시트 로더.

사용처:
- parser.py: 워크북01 추출 stories의 examCode → "올포 N강 본문M" 헤더 변환
- vocab_fetcher.py: Vol1 중요어휘에 올포 본문/요약 인입

xlsx 출처: ~/.../EBS 작업/올림포스 고2 기출/고2 올림포스 기출 지문 출처 2026 편집.xlsx
시트 시트 26올포기출: 1tDHMuLfHHpPUQpVnQGI2facI66pqbfdmMt5kWt59x8c gid=1097793467
"""

import os
import re
import csv
import io
import urllib.request
import unicodedata


def _nfc(s):
    return unicodedata.normalize("NFC", s) if s else s


# ──────────────────────────────────────────────────────────────────────
# 1) xlsx → {YYMMQQ: (lecture, body_label)}
# ──────────────────────────────────────────────────────────────────────
def load_olpo_lecture_map(xlsx_path, year='2026', log=None):
    """고2 올림포스 기출 출처 xlsx를 읽어 매핑 빌드.

    매핑 형식:
        {YYMMQQ: (lecture_int, body_label)}
        - lecture_int: 1~18 (xlsx 헤더 행에서 추출)
        - body_label: "분석" (학평Analysis 행) 또는 "본문N" (1, 2, 3, ... N번 행)

    빨간색 폰트 셀은 무시 (사용자 지시: "아래 빨간색 폰트는 무시").

    Returns:
        dict {str(YYMMQQ): (int, str)}
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        if log:
            log("[OLPO MAP] openpyxl 미설치 — 매핑 비활성")
        return {}

    if not os.path.exists(xlsx_path):
        if log:
            log(f"[OLPO MAP] xlsx 없음: {xlsx_path}")
        return {}

    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        if log:
            log(f"[OLPO MAP] 로드 실패: {e}")
        return {}

    if year not in wb.sheetnames:
        if log:
            log(f"[OLPO MAP] 시트 없음: {year} (있는 시트: {wb.sheetnames})")
        return {}

    ws = wb[year]
    mapping = {}

    # Row 1: header — 1강(목적), 2강(심경), 3강(주장), ...
    headers = []
    for cell in ws[1]:
        v = cell.value
        if v and "강" in str(v):
            m = re.match(r"^\s*(\d+)\s*강", str(v))
            headers.append(int(m.group(1)) if m else None)
        else:
            headers.append(None)

    # Row 2+: data
    # Row 2: "학평Analysis" → label "분석"
    # Row 3+: 번호 (1, 2, 3, ...) → label "본문N"
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        # row[0] = label cell (학평Analysis or 번호)
        label_cell = row[0]
        label_val = label_cell.value
        if label_val is None:
            continue

        # 본문 라벨 결정
        if isinstance(label_val, str) and "Analysis" in label_val:
            body_label = "분석"
        elif isinstance(label_val, int):
            body_label = f"본문{label_val}"
        elif isinstance(label_val, str) and label_val.strip().isdigit():
            body_label = f"본문{int(label_val.strip())}"
        else:
            continue

        for col_idx, cell in enumerate(row):
            if col_idx == 0:
                continue
            if col_idx >= len(headers):
                continue
            lec = headers[col_idx]
            if lec is None:
                continue
            code = cell.value
            if code is None:
                continue

            # 빨간색 폰트 무시 (사용자 지시)
            try:
                font = cell.font
                if font and font.color and font.color.rgb:
                    rgb = str(font.color.rgb).upper()
                    # FFFF0000 = pure red, also FF... pattern
                    if rgb.endswith("FF0000") or rgb in ("FFFF0000",):
                        continue
            except Exception:
                pass

            # YYMMQQ는 6자리 숫자
            code_str = str(code).strip()
            if not re.match(r"^\d{6}$", code_str):
                continue

            mapping[code_str] = (lec, body_label)

    if log:
        log(f"[OLPO MAP] {year} 시트에서 {len(mapping)}개 코드 매핑 로드")

    return mapping


# ──────────────────────────────────────────────────────────────────────
# 2) 26올포기출 시트 → {(lecture, body_label): {body, summary}}
# ──────────────────────────────────────────────────────────────────────
SHEET_ID = "1tDHMuLfHHpPUQpVnQGI2facI66pqbfdmMt5kWt59x8c"
DEFAULT_GIDS = {
    "26올포기출": "1097793467",
}


def _fetch_csv(sheet_id, gid):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    data = urllib.request.urlopen(req, timeout=30).read().decode("utf-8")
    return list(csv.reader(io.StringIO(data)))


def _convert_smart_quotes(text):
    """ASCII straight quotes → typographic(smart) quotes.

    템플릿 단락 스타일 "2. 중요어휘(왼)" / "2. 중요어휘_원문(오)"의 GREP 스타일이
    `^"` 와 `"$` 패턴(ASCII 따옴표)을 투명 처리하므로 문단 앞뒤의 따옴표가 사라짐.
    스마트 따옴표(U+201C/201D, U+2018/2019)로 변환하면 GREP 매칭을 회피하고
    인쇄 품질도 높아짐.

    교차 변환: " → " ↔ " 과 ' → ' ↔ ' 형태로 alternating.
    """
    if not text:
        return text
    out = []
    in_double = False
    in_single = False
    prev_ch = ""
    for ch in text:
        if ch == '"':
            # 단어 경계 휴리스틱: 직전이 알파벳/숫자면 close, 아니면 open
            if in_double or (prev_ch and (prev_ch.isalnum() or prev_ch in ".,;:!?")):
                out.append("”")  # ”
                in_double = False
            else:
                out.append("“")  # “
                in_double = True
        elif ch == "'":
            # 영문 축약형(it's, don't 등)에서 알파벳 사이에 있으면 그대로 typographic ’
            if prev_ch and prev_ch.isalpha():
                out.append("’")  # ’
            elif in_single or (prev_ch and (prev_ch.isalnum() or prev_ch in ".,;:!?")):
                out.append("’")
                in_single = False
            else:
                out.append("‘")  # ‘
                in_single = True
        else:
            out.append(ch)
        prev_ch = ch
    return "".join(out)


def load_olpo_vocab_sheet(year_key="26올포기출", log=None):
    """올포 기출 중요어휘 시트를 읽어 (lecture, body_label) → {body, summary} 매핑 빌드.

    시트 구조:
        Row 0: ['1강', '1번', '2번', ..., '16번']  (17 cols)
        Row 1: data for 1강 (col 0 비어있음, col 1=1번, col 2=2번, ...)
        Row 2: ['2강', '1번', '2번', ...]
        Row 3: data for 2강
        ...

    각 셀은 본문 + ">>   내용요약   " + 요약 형식.

    Returns:
        dict {(int_lecture, str_body_label): {"body": ..., "summary": ...}}
    """
    gid = DEFAULT_GIDS.get(year_key)
    if not gid:
        if log:
            log(f"[OLPO VOCAB] 알 수 없는 시트 키: {year_key}")
        return {}

    try:
        rows = _fetch_csv(SHEET_ID, gid)
    except Exception as e:
        if log:
            log(f"[OLPO VOCAB] CSV 로드 실패: {e}")
        return {}

    if not rows:
        return {}

    mapping = {}
    current_lec = None
    current_headers = []  # ["1번", "2번", ...]

    for row in rows:
        if not row:
            continue
        first = (row[0] or "").strip()
        # 강 헤더 행
        m_lec = re.match(r"^\s*(\d+)\s*강", first)
        if m_lec:
            current_lec = int(m_lec.group(1))
            current_headers = [(c or "").strip() for c in row[1:]]
            continue

        # 데이터 행 (col 0 비어있음)
        if first or current_lec is None:
            continue

        for idx, cell in enumerate(row[1:]):
            cell = (cell or "").strip()
            if not cell or idx >= len(current_headers):
                continue
            header = current_headers[idx]
            m_num = re.match(r"^\s*(\d+)\s*번", header)
            if not m_num:
                continue
            num = int(m_num.group(1))
            body_label = f"본문{num}"

            # 본문 / 내용요약 분리
            split_marker = re.search(r">>\s*내용요약\s*", cell)
            if split_marker:
                body = cell[: split_marker.start()].strip()
                summary = cell[split_marker.end() :].strip()
            else:
                body = cell
                summary = ""

            # ASCII 따옴표 → 스마트 따옴표 변환
            # (템플릿 GREP 스타일 `^"` `"$` 가 ASCII 인용부호를 투명 처리하는 문제 회피)
            body = _convert_smart_quotes(body)
            summary = _convert_smart_quotes(summary)

            mapping[(current_lec, body_label)] = {"body": body, "summary": summary}

    if log:
        log(f"[OLPO VOCAB] {year_key} 시트에서 {len(mapping)}개 (강,본문) 매핑 로드")

    return mapping


# ──────────────────────────────────────────────────────────────────────
# 3) 통합 헬퍼
# ──────────────────────────────────────────────────────────────────────
def code_to_olpo_header(code, lecture_map):
    """YYMMQQ 코드 → "올포 N강 본문M" 형식 헤더. 매핑 없으면 None."""
    if not lecture_map:
        return None
    info = lecture_map.get(str(code).strip())
    if not info:
        return None
    lec, body_label = info
    return f"올포 {lec}강 {body_label}"


def find_olpo_xlsx_path(content_base):
    """내신대비 폴더에서 고2 올림포스 기출 출처 xlsx 경로 찾기.

    경로: EBS 작업/올림포스 고2 기출/고2 올림포스 기출 지문 출처 *.xlsx
    """
    base = os.path.join(content_base, "EBS 작업", "올림포스 고2 기출")
    # 미러 redirect — Drive zombie 시 GDrive listdir EDEADLK 회피
    from engines.final.extractor import _resolve_local_mirror as _rlm
    base = _rlm(base, None)
    if not os.path.isdir(base):
        return None
    for f in os.listdir(base):
        fn = _nfc(f)
        if (
            fn.endswith(".xlsx")
            and "출처" in fn
            and "고2" in fn
            and not f.startswith("~")
        ):
            return os.path.join(base, f)
    return None
