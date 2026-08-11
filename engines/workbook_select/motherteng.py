"""마더텅 지문 출처 엑셀 → 코드(YYMMQQ) ↔ 마더텅 지문번호 매핑.

엑셀 구조 (2026 마더텅 고1 영어 독해 지문 출처.xlsx, Sheet1):
  - 헤더 행: 강 이름들이 각각 하나의 열 ('03 필자의 주장', '11 빈칸 추론' …)
  - A열: 그 강 안에서의 마더텅 지문번호 (1, 2, 3 …)
  - 셀: 모의고사 코드 YYMMQQ 6자리

같은 강이 공백 유무만 다른 이름으로 두 열 존재할 수 있다
('11 빈칸 추론' 75개 = 전체, '11 빈칸추론' 20개 = 실제 조판 대상 부분집합).
번호의 기준은 **전체 열**이므로, 같은 강에 여러 열이 있으면 항목이 많은 쪽을
기준으로 삼는다 (부분집합 열의 행 위치는 전체 열과 같아 결과는 동일하지만,
부분집합만 있는 경우에도 안전하게 동작한다).

사용자 확정 스펙(2026-08-07): 지문번호는 이 행번호를 그대로 쓴다. 짝수만
뽑힌 워크북이면 지문번호도 2, 4, 6 … 이며 1,2,3,4 로 다시 매기지 않는다.
"""
import os
import re
import unicodedata

_RE_CODE = re.compile(r"^\d{6}$")
_RE_GANG = re.compile(r"^\s*(\d{1,2})\s*(.*)$")


def _nfc(s):
    return unicodedata.normalize("NFC", str(s or "")).strip()


def load_index(xlsx_path, header_row=3, first_col=1):
    """엑셀 → {code: {"gang": int, "gang_name": str, "no": int}}

    강이 여러 열로 쪼개져 있으면 항목 수가 많은 열을 채택한다.
    파일이 없거나 openpyxl 이 없으면 빈 dict (호출부가 폴백)."""
    path = os.path.expanduser(xlsx_path or "")
    if not path or not os.path.exists(path):
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return {}
    sheet = wb[wb.sheetnames[0]]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    if len(rows) < header_row:
        return {}
    header = rows[header_row - 1]

    # 열 인덱스 → (강번호, 강이름, {code: no})
    cols = []
    for ci in range(first_col, len(header)):
        name = _nfc(header[ci])
        if not name:
            continue
        m = _RE_GANG.match(name)
        if not m:
            continue
        gang = int(m.group(1))
        entries = {}
        for r in rows[header_row:]:
            if ci >= len(r):
                continue
            code = _nfc(r[ci])
            if not _RE_CODE.match(code):
                continue
            try:
                no = int(_nfc(r[0]))
            except (TypeError, ValueError):
                continue
            entries[code] = no
        if entries:
            cols.append((gang, name, entries))

    # 같은 강 여러 열 → 항목 많은 쪽 채택
    best = {}
    for gang, name, entries in cols:
        if gang not in best or len(entries) > len(best[gang][1]):
            best[gang] = (name, entries)

    index = {}
    for gang, (name, entries) in best.items():
        for code, no in entries.items():
            # 같은 코드가 여러 강에 있으면 먼저 만난 강 유지 (강 지정 시 필터)
            index.setdefault(code, {"gang": gang, "gang_name": name, "no": no})
    return index


def load_columns(xlsx_path, header_row=3, first_col=1):
    """엑셀 → [(강번호, 열이름, [(no, code), ...])] — 열을 합치지 않고 그대로.

    load_index 는 강별로 항목이 많은 열 하나만 남기지만, **조판 대상 선택분은
    항목이 적은 쪽 열**인 경우가 있다(실측: '11 빈칸 추론' 75개=전체,
    '11 빈칸추론' 20개=이 워크북의 선택분). 강 번호만으로 코드를 뽑으려면
    두 열을 다 봐야 한다.
    """
    path = os.path.expanduser(xlsx_path or "")
    if not path or not os.path.exists(path):
        return []
    try:
        import openpyxl
    except ImportError:
        return []
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []
    sheet = wb[wb.sheetnames[0]]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    if len(rows) < header_row:
        return []
    header = rows[header_row - 1]

    cols = []
    for ci in range(first_col, len(header)):
        name = _nfc(header[ci])
        if not name:
            continue
        m = _RE_GANG.match(name)
        if not m:
            continue
        entries = []
        for r in rows[header_row:]:
            if ci >= len(r):
                continue
            code = _nfc(r[ci])
            if not _RE_CODE.match(code):
                continue
            try:
                no = int(_nfc(r[0]))
            except (TypeError, ValueError):
                continue
            entries.append((no, code))
        if entries:
            cols.append((int(m.group(1)), name, entries))
    return cols


def codes_for_gang(xlsx_path, gang, log=None):
    """강 번호만으로 조판 대상 코드를 뽑는다 → (codes, numbers, 채택열이름).

    codes  : 마더텅 번호 오름차순 코드 리스트 (그대로 조판 순서)
    numbers: {code: 마더텅 지문번호}

    선택 규칙 — 같은 강에 열이 여럿이면 **항목이 적은 쪽(부분집합)** 이
    이 워크북의 선택분이다. 번호는 어느 열에서 읽어도 같은 행 기준이라
    동일하다. 열이 하나뿐이면 그 열 전체가 대상.
    침묵 금지: 어느 열을 왜 골랐는지 로그로 남긴다.
    """
    cols = [c for c in load_columns(xlsx_path) if c[0] == int(gang)]
    if not cols:
        if log:
            log("[마더텅] 엑셀에 %s강 열 없음" % gang)
        return [], {}, ""
    cols.sort(key=lambda c: len(c[2]))
    gang_no, name, entries = cols[0]
    if log and len(cols) > 1:
        log("[마더텅] %d강 열 %d개 — '%s'(%d개) 채택 / 후보: %s"
            % (gang_no, len(cols), name, len(entries),
               ", ".join("'%s'(%d)" % (c[1], len(c[2])) for c in cols)))
    elif log:
        log("[마더텅] %d강 '%s' %d개" % (gang_no, name, len(entries)))
    entries = sorted(entries, key=lambda t: t[0])
    codes, numbers, seen = [], {}, set()
    for no, code in entries:
        if code in seen:
            continue
        seen.add(code)
        codes.append(code)
        numbers[code] = no
    return codes, numbers, name


def numbers_for(index, codes, gang=None):
    """코드 리스트 → {code: 마더텅 번호}. gang 지정 시 그 강만 인정."""
    out = {}
    for c in codes:
        rec = index.get(c)
        if not rec:
            continue
        if gang is not None and rec["gang"] != gang:
            continue
        out[c] = rec["no"]
    return out


def sort_by_number(codes, numbers):
    """마더텅 번호 순으로 정렬. 번호 없는 코드는 뒤에 원래 순서로."""
    known = [c for c in codes if c in numbers]
    unknown = [c for c in codes if c not in numbers]
    known.sort(key=lambda c: numbers[c])
    return known + unknown
