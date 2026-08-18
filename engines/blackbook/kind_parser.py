"""
블랙북 신구조(PART1·PART2 + 어휘) 다종 파서

5종 원고(part1 문제 / part1 어휘 / part2 문제 / part1 해설 / part2 해설)를
docx 또는 hwpx 어느 쪽이든 받아 InDesign 단락 배열로 변환.

출력 단락 구조:
    {"text": str, "type": str, "runs": [{"t": str, "ul": bool}, ...]}

신규 단락 타입 (어휘):
    vocab_header    : "<3과 본문1>" 등 본문 묶음 헤더
    vocab_passage   : 어휘 학습용 본문 (Host/Dr ...)
    vocab_category  : "〈연관 어휘〉" 라벨
    vocab_item      : 표제어+뜻+동의어 한 덩어리
"""

import os
import re
import unicodedata
import zipfile
from typing import List, Dict, Optional

from engines.blackbook.parser import (
    _classify,
    _parse_answer_lines,
    _parse_paragraphs,
    _extract_hwpx_answer_lines,
)

try:
    from docx import Document
    _DOCX_OK = True
except ImportError:
    _DOCX_OK = False


# ── 파일명 분류 ─────────────────────────────────────────────────────────────

_KIND_KEYS = (
    "part1_problem", "part1_vocab", "part2_problem",
    "part1_answer",  "part2_answer",
)


def classify_filename(fname: str) -> Optional[str]:
    """파일명 키워드로 5종 중 하나 분류. 매칭 실패 시 None.

    Part 번호 추출 우선순위 (호번호 '01', '02' 등의 false-match 회피):
        1. 'Part'/'PART'/'part' 토큰 바로 뒤의 숫자 (예: 'Part2')
        2. 'N부' 한글 표기 (예: '1부')
    """
    nfc = unicodedata.normalize("NFC", fname)
    norm = re.sub(r"\s+", "", nfc)
    norm_low = norm.lower()

    # part 번호 추출 (호번호 '01' 등의 false-match 회피)
    part_num = None
    m = re.search(r"part\s*(\d+)", norm_low)
    if m:
        part_num = int(m.group(1))
    else:
        # 'N부' — 앞 글자가 숫자가 아닌 경우만 (호번호 뒤 직접 결합 제외)
        m = re.search(r"(?<!\d)(\d+)부", nfc)
        if m:
            part_num = int(m.group(1))

    is_vocab   = ("어휘" in nfc)
    is_answer  = ("해설" in nfc) or ("정답" in nfc)
    is_problem = ("문제" in nfc) and not is_vocab and not is_answer

    if part_num == 1 and is_vocab:   return "part1_vocab"
    if part_num == 1 and is_answer:  return "part1_answer"
    if part_num == 2 and is_answer:  return "part2_answer"
    if part_num == 1 and is_problem: return "part1_problem"
    if part_num == 2 and is_problem: return "part2_problem"

    # part 표기가 빠진 어휘 파일 → part1 (현재 part1만 어휘 있음)
    if is_vocab and part_num != 2:
        return "part1_vocab"

    return None


def classify_filename_reason(fname: str) -> str:
    """분류 실패(None) 가시화용 근거 문자열 (배치 G — 판정 로직 무변경).

    classify_filename 과 같은 신호를 재평가해 '무엇을 읽었고 무엇이 빠졌는지'
    를 만든다. 실패 로그에 붙여 반침묵(경고 후 무시)을 진단 가능하게 한다.
    """
    try:
        nfc = unicodedata.normalize("NFC", fname)
        norm_low = re.sub(r"\s+", "", nfc).lower()
        part_num = None
        m = re.search(r"part\s*(\d+)", norm_low)
        if m:
            part_num = int(m.group(1))
        else:
            m = re.search(r"(?<!\d)(\d+)부", nfc)
            if m:
                part_num = int(m.group(1))
        signals = []
        signals.append(f"Part번호={'미검출' if part_num is None else part_num}")
        signals.append("어휘" if "어휘" in nfc else "어휘토큰없음")
        signals.append("정답/해설" if ("해설" in nfc or "정답" in nfc) else "정답토큰없음")
        signals.append("문제" if "문제" in nfc else "문제토큰없음")
        return ", ".join(signals) + " — 5종(part1 문제/어휘/해설, part2 문제/해설) 어디에도 미대응"
    except Exception:
        return "근거 평가 실패"


# ── docx → 단락 (원시) ─────────────────────────────────────────────────────

def _docx_paragraphs(docx_path: str) -> List[Dict]:
    """docx → [{"text": str, "runs": [{"t": str, "ul": bool}]}]. 빈 단락 제외."""
    if not _DOCX_OK:
        raise RuntimeError("python-docx 미설치. pip install python-docx 실행 후 재시도.")
    doc = Document(docx_path)
    out = []
    for p in doc.paragraphs:
        text = p.text
        if not text or not text.strip():
            continue
        runs = []
        for r in p.runs:
            if r.text:
                runs.append({"t": r.text, "ul": bool(r.underline)})
        out.append({"text": text, "runs": runs})
    return out


# ── hwpx → 단락 (원시) ─────────────────────────────────────────────────────

def _hwpx_paragraphs(hwpx_path: str) -> List[Dict]:
    """hwpx → 단락 목록. parser.py의 헬퍼 재활용."""
    with zipfile.ZipFile(hwpx_path, "r") as zf:
        names = zf.namelist()
        sec_name = next(
            (n for n in names if re.search(r"section\d+\.xml", n, re.I)), None
        )
        if not sec_name:
            raise ValueError(f"section XML 없음: {hwpx_path}")
        section_xml = zf.read(sec_name).decode("utf-8")
        header_name = next(
            (n for n in names if re.search(r"header\.xml", n, re.I)), None
        )
        header_xml = zf.read(header_name).decode("utf-8") if header_name else ""
    return _parse_paragraphs(section_xml, header_xml)


def _raw_paragraphs(path: str) -> List[Dict]:
    """확장자 자동 분기."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _docx_paragraphs(path)
    if ext in (".hwpx",):
        return _hwpx_paragraphs(path)
    # hwp 바이너리는 본문 파서 미지원 (정답 파일만 hwp5로 지원)
    raise ValueError(f"지원하지 않는 확장자: {ext} ({os.path.basename(path)})")


# ── 문제 파일 헤더 패턴 (조판 제외) ────────────────────────────────────────

_HEADER_PATTERNS = (
    re.compile(r"^블랙북\d+\s*[—–\-]"),
    re.compile(r".+고등학교\s+\d+학년\s+영어"),
    re.compile(r"^\d+유형\s*[×x]\s*\d+문항"),
)


def _is_problem_header(text: str) -> bool:
    t = text.strip()
    for p in _HEADER_PATTERNS:
        if p.match(t):
            return True
    return False


# ── 문제 파일 파싱 ─────────────────────────────────────────────────────────

_RE_PROBLEM_NUM = re.compile(r"^\s*(\d+)\s*\.")


def parse_problem(path: str) -> List[Dict]:
    """문제 docx/hwpx → 문제 단위 그룹 배열.

    출력 구조 (frame당 한 문제 매핑용):
        [
          {"problem_num": 1, "paragraphs": [{type, text, runs}, ...]},
          {"problem_num": 2, "paragraphs": [...]},
          ...
        ]

    'N.' 패턴 단락(type='problem')이 새 문제의 시작. 그 사이 단락(지문/선지/조건 등)은
    같은 문제 그룹에 묶임. 상단 헤더 3줄은 제외.
    """
    paras = _raw_paragraphs(path)
    groups = []
    current = None
    in_condition = False
    header_window = True

    for p in paras:
        text = p["text"]
        if header_window:
            if _is_problem_header(text):
                continue
            if re.match(r"^\d+\.", text.strip()):
                header_window = False

        typ = _classify(text, in_condition)
        if typ in ("blank", "day_header"):
            continue

        if typ == "condition":
            in_condition = True
        elif typ in ("problem", "passage_header", "next_question",
                     "choice", "answer_blank"):
            in_condition = False

        para_obj = {"text": text, "type": typ, "runs": p.get("runs", [])}

        if typ == "problem":
            # 새 문제 그룹 시작
            if current:
                groups.append(current)
            m = _RE_PROBLEM_NUM.match(text)
            current = {
                "problem_num": int(m.group(1)) if m else len(groups) + 1,
                "paragraphs": [para_obj],
            }
        else:
            # 첫 problem 등장 전 단락은 버림 (헤더 잔여물 방지)
            if current is None:
                continue
            current["paragraphs"].append(para_obj)

    if current:
        groups.append(current)

    # 문장 삽입 / 요약 문제는 특정 단락에 'background' 타입 부여
    _mark_background_paragraphs(groups)
    return groups


# 문제 단락 텍스트 패턴: "주어진 글/문장" + "이어질/들어가기" → 첫 passage 배경 대상
_RE_NEEDS_GIVEN_FIRST = re.compile(r"주어진\s*(문장|글)")
# 단락 텍스트가 [요약문] 라벨 자체인지
_RE_SUMMARY_LABEL = re.compile(r"^\s*[\[<〈【]\s*요약\s*문?\s*[\]>〉】]\s*$")


def _mark_background_paragraphs(groups):
    """문제 구조를 읽어 배경 단락 식별.

    1) 문제 텍스트에 '주어진 글/문장' 포함 (글의 순서 / 문장 삽입):
       → 문제 직후 첫 passage 단락 = 주어진 단락 = background
    2) 단락 중 [요약문] 라벨 발견:
       → 라벨 자체 + 그 다음 단락 모두 background
    """
    for grp in groups:
        paras = grp.get("paragraphs", [])
        if not paras:
            continue

        # 1) 글의 순서 / 문장 삽입 — 첫 passage에 background
        prob_text = paras[0]["text"]
        if _RE_NEEDS_GIVEN_FIRST.search(prob_text):
            for i in range(1, len(paras)):
                if paras[i]["type"] == "passage":
                    paras[i]["type"] = "background"
                    break

        # 2) [요약문] 라벨 단락 식별 → 라벨 + 다음 단락 background
        for i in range(len(paras)):
            if _RE_SUMMARY_LABEL.match(paras[i]["text"].strip()):
                paras[i]["type"] = "background"
                # 라벨 다음 단락(요약문 본체)도 background
                if i + 1 < len(paras) and paras[i+1]["type"] in ("passage", "background"):
                    paras[i+1]["type"] = "background"
                break


# ── 어휘 파일 파싱 ─────────────────────────────────────────────────────────

_RE_VOCAB_HEADER   = re.compile(r"^\s*<\s*\d+과\s*본문\s*\d+\s*>")
_RE_VOCAB_CATEGORY = re.compile(r"^\s*〈\s*연관\s*어휘\s*〉")


def parse_vocab(path: str) -> List[Dict]:
    """어휘 docx/hwpx → 본문 단위 그룹 배열 (frame당 한 본문 묶음 매핑용).

    출력 구조:
        [
          {"header": "<3과 본문1>", "paragraphs": [{type, text, runs}, ...]},
          {"header": "<3과 본문2>", "paragraphs": [...]},
          ...
        ]

    단락 타입 분류:
        <N과 본문M>     → vocab_header
        〈연관 어휘〉   → vocab_category
        그 이전 단락    → vocab_passage  (어휘 학습용 본문)
        그 이후 단락    → vocab_item     (표제어+뜻+동의어)
    """
    paras = _raw_paragraphs(path)
    out = []  # 호환성 위해 보존 (단락 평탄 배열)
    groups = []
    current = None
    in_items = False
    for p in paras:
        text = p["text"]
        runs = p.get("runs", [])
        if _RE_VOCAB_HEADER.match(text):
            typ = "vocab_header"
            in_items = False
        elif _RE_VOCAB_CATEGORY.match(text):
            typ = "vocab_category"
            in_items = True
        elif in_items:
            typ = "vocab_item"
        else:
            typ = "vocab_passage"

        para_obj = {"text": text, "type": typ, "runs": runs}

        if typ == "vocab_header":
            if current:
                groups.append(current)
            current = {"header": text, "paragraphs": [para_obj]}
        else:
            if current is None:
                # 첫 vocab_header 이전 단락은 기본 그룹으로
                current = {"header": "", "paragraphs": []}
            current["paragraphs"].append(para_obj)

    if current:
        groups.append(current)
    return groups


# ── 정답/해설 파일 파싱 ────────────────────────────────────────────────────

def parse_answers(path: str) -> List[Dict]:
    """정답/해설 docx/hwpx/hwp → [{"num": int, "answer": str, "explanation": str}]."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        if not _DOCX_OK:
            raise RuntimeError("python-docx 미설치.")
        doc = Document(path)
        lines = []
        for p in doc.paragraphs:
            for sub in p.text.split("\n"):
                sub = sub.strip()
                if sub:
                    lines.append(sub)
        return _parse_answer_lines(lines)

    if ext == ".hwpx":
        lines = _extract_hwpx_answer_lines(path)
        return _parse_answer_lines(lines)

    if ext == ".hwp":
        # parser.py의 parse_answers_hwp가 hwp 바이너리도 지원
        from engines.blackbook.parser import parse_answers_hwp
        return parse_answers_hwp(path)

    raise ValueError(f"지원하지 않는 확장자: {ext} ({os.path.basename(path)})")


# ── 디스패처 ────────────────────────────────────────────────────────────────

def parse_by_kind(kind: str, path: str) -> List[Dict]:
    """kind에 맞는 파서 호출. 알 수 없는 kind면 ValueError."""
    if kind in ("part1_problem", "part2_problem"):
        return parse_problem(path)
    if kind == "part1_vocab":
        return parse_vocab(path)
    if kind in ("part1_answer", "part2_answer"):
        return parse_answers(path)
    raise ValueError(f"알 수 없는 kind: {kind}")


# ── 구조 기반 분류 (파일명 분류 실패 시 폴백) ──────────────────────────────
#
# 파일명 규칙(`Part1 문제`)은 호마다 바뀐다 — `블랙북01_문제_교과서_1과.docx`
# 처럼 Part 표기가 없는 원고가 통째로 버려지던 문제(W38) 대응.
# 파일명으로 판정되면 그 결과를 그대로 쓰고, 실패했을 때만 원고 구조를 읽는다.

_RE_ANSWER_MARK = re.compile(r"^\s*\[\s*정답\s*\]")
_RE_EXPL_MARK   = re.compile(r"^\s*\[\s*해설\s*\]")
_RE_CHOICE_MARK = re.compile(r"^\s*[①-⑳]")   # ①~⑳
_RE_TYPE_HEADER = re.compile(r"^\s*\d+\s*유형\s*[×x*]\s*\d+\s*문항")
_RE_LESSON_KO   = re.compile(r"(\d+)\s*(?:과|강)")
_RE_LESSON_EN   = re.compile(r"\bLesson\s*(\d+)", re.I)


def _sniff_paragraphs(path: str, limit: int = 400) -> List[str]:
    """원고 앞부분 단락 텍스트 (구조 판정용). 실패하면 빈 목록."""
    try:
        paras = _raw_paragraphs(path)
    except Exception:
        return []
    return [unicodedata.normalize("NFC", p.get("text", "")) for p in paras[:limit]]


def classify_content(path: str, texts: Optional[List[str]] = None) -> Optional[str]:
    """원고 구조로 5종 분류. 판정 불가면 None.

    앵커 (형식이 아니라 구조):
        어휘 — `<N과 본문M>` 헤더 + `〈연관 어휘〉` 카테고리
        해설 — `[정답]` 줄 반복
        문제 — `N.` 발문 + `①`~ 선지 반복 (`N유형 × M문항` 헤더면 part2)
    """
    if texts is None:
        texts = _sniff_paragraphs(path)
    if not texts:
        return None

    vocab_hdr = sum(1 for t in texts if _RE_VOCAB_HEADER.match(t))
    vocab_cat = sum(1 for t in texts if _RE_VOCAB_CATEGORY.match(t))
    answers   = sum(1 for t in texts if _RE_ANSWER_MARK.match(t))
    expls     = sum(1 for t in texts if _RE_EXPL_MARK.match(t))
    problems  = sum(1 for t in texts if _RE_PROBLEM_NUM.match(t))
    choices   = sum(1 for t in texts if _RE_CHOICE_MARK.match(t))
    is_type   = any(_RE_TYPE_HEADER.match(t) for t in texts[:5])

    if vocab_cat >= 1 and (vocab_hdr >= 1 or vocab_cat >= 2):
        return "part1_vocab"
    if answers >= 2 or (answers >= 1 and expls >= 1):
        return "part2_answer" if is_type else "part1_answer"
    if problems >= 2 and choices >= 2:
        return "part2_problem" if is_type else "part1_problem"
    return None


def has_part_token(fname: str) -> bool:
    """파일명에 'Part1' / '2부' 같은 Part 표기가 실제로 있는지."""
    nfc = unicodedata.normalize("NFC", fname)
    norm_low = re.sub(r"\s+", "", nfc).lower()
    if re.search(r"part\s*\d+", norm_low):
        return True
    return bool(re.search(r"(?<!\d)(\d+)부", nfc))


def classify_file(path: str, texts: Optional[List[str]] = None) -> Optional[str]:
    """파일명 우선, 실패 시 원고 구조로 분류."""
    name = os.path.basename(path)
    for prefix in ("manuscript_", "answers_"):
        if name.startswith(prefix):
            name = name[len(prefix):]
            break
    kind = classify_filename(name)
    if kind:
        return kind
    return classify_content(path, texts)


def lesson_of(path: str, texts: Optional[List[str]] = None) -> Optional[int]:
    """원고가 다루는 과(강) 번호. 파일명 → 원고 헤더 순으로 찾는다."""
    name = unicodedata.normalize("NFC", os.path.basename(path))
    stem = os.path.splitext(name)[0]
    # 파일명의 '블랙북01' 같은 호번호는 과 표기가 아니므로 '과/강' 접미만 본다
    m = _RE_LESSON_KO.search(stem)
    if m:
        return int(m.group(1))
    if texts is None:
        texts = _sniff_paragraphs(path, limit=60)
    for t in texts[:60]:
        m = _RE_VOCAB_HEADER.match(t) or None
        if m:
            m2 = _RE_LESSON_KO.search(t)
            if m2:
                return int(m2.group(1))
        m3 = _RE_LESSON_EN.search(t)
        if m3:
            return int(m3.group(1))
        m4 = _RE_LESSON_KO.search(t)
        if m4:
            return int(m4.group(1))
    return None


def book_title_of(texts: Optional[List[str]]) -> Optional[str]:
    """원고 첫 줄에서 교재명만 추출.

        '공통영어2 능률(오)  Lesson 1  A Band United by a Name'
            → '공통영어2 능률(오)'
    """
    if not texts:
        return None
    for t in texts[:5]:
        t = (t or "").strip()
        if not t or _RE_VOCAB_HEADER.match(t) or _RE_VOCAB_CATEGORY.match(t):
            continue    # 어휘 원고 헤더(<N과 본문M>)는 교재명이 아니다
        m = _RE_LESSON_EN.search(t)
        if m:
            head = t[:m.start()].strip()
            return head or None
        m2 = _RE_LESSON_KO.search(t)
        if m2 and m2.start() > 0:
            head = t[:m2.start()].strip()
            if len(head) >= 2:
                return head
        return None
    return None


def subtitle_of(book: Optional[str], lesson_num: Optional[int]) -> str:
    """러닝 헤더용 부제목 문자열. 예: '공통영어2 능률(오)  1과'."""
    label = f"{lesson_num}과" if lesson_num is not None else ""
    if book and label:
        return f"{book}  {label}"
    return book or label
